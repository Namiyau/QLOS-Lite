from __future__ import annotations

import json
import os
import sys
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from qlos_lite.config import load_qlos_lite_config  # noqa: E402
from qlos_lite.onebot_client import OneBotClient  # noqa: E402


SOURCES = [
    {
        "name": "Open WebUI Features",
        "url": "https://docs.openwebui.com/features/",
        "takeaway": "工具、函数、OpenAPI、MCP、技能、提示词模板可作为扩展入口。",
    },
    {
        "name": "Open WebUI Pipelines",
        "url": "https://docs.openwebui.com/features/extensibility/pipelines/",
        "takeaway": "消息前后处理、路由、过滤、限流、监控都可放到网关层或外部服务。",
    },
    {
        "name": "LibreChat Resumable Streams",
        "url": "https://www.librechat.ai/docs/features/resumable_streams",
        "takeaway": "断线续流、多端同步、后台生成，适合 Hermi 的桌面和手机端。",
    },
    {
        "name": "LibreChat User Memory",
        "url": "https://www.librechat.ai/docs/features/memory",
        "takeaway": "结构化 key/value 记忆、用户可编辑、token 限额，适合记忆审计 inbox。",
    },
    {
        "name": "Dify Human Input API Integration Flow",
        "url": "https://docs.dify.ai/en/cloud/use-dify/nodes/hitl-api-integration-flow",
        "takeaway": "工作流可暂停、发出 human_input_required 事件，再由外部客户端提交审批。",
    },
    {
        "name": "Dify Human Input",
        "url": "https://docs.dify.ai/en/cloud/use-dify/nodes/human-input",
        "takeaway": "审批表单可包含说明、变量、操作按钮、超时策略。",
    },
    {
        "name": "AnythingLLM Docs",
        "url": "https://docs.useanything.com/",
        "takeaway": "AI agents、日志、API、权限、安全、个性化记忆、模型路由都是个人 AI 客户端常见模块。",
    },
    {
        "name": "AnythingLLM Scheduled Jobs",
        "url": "https://docs.anythingllm.com/scheduled-jobs/overview",
        "takeaway": "定时任务应保留完整 trace、产物、工具调用记录，便于事后审计。",
    },
]


QLOS_HERMI_CONTEXT = """
QLOS-Lite 是 QQ / NapCat / OneBot 到 Hermes 的薄安全社交网关。它负责收 QQ 消息、识别 owner / non-owner、写身份文件、做风险拦截、处理附件、持久队列、trace/doctor、调用 Hermes API、清洗并拆分 QQ 输出、把回复发回 QQ。

Hermes 是主脑，负责人格、记忆、工具、文件分析、沙箱任务、长期策略。QLOS-Lite 不替代 Hermes，只做来源、身份、权限、防越权、日志和传输。

Hermi 是未来自建桌面端 / 手机端 / PWA。它会像 ChatGPT 一样有会话列表，并和 QQ 私聊、群聊、Hermi 自建会话共享同一套 Hermes API 和记忆。它还要有 owner 控制台、审批 inbox、记忆审计、额度限制、朋友访问、图片/语音、多端同步、远程遥控和主动消息入口。
"""


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def build_codex_doc(path: Path) -> None:
    doc = Document()
    doc.add_heading("codex思索：QLOS-Lite 与 Hermi 可玩功能", 0)
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    add_heading(doc, "项目定位", 1)
    doc.add_paragraph(QLOS_HERMI_CONTEXT.strip())

    add_heading(doc, "调研启发", 1)
    add_bullets(doc, [f"{s['name']}：{s['takeaway']}（{s['url']}）" for s in SOURCES])

    add_heading(doc, "近期最值得做", 1)
    add_numbered(
        doc,
        [
            "QLOS Doctor Pro：一键检查 NapCat、QLOS、Hermes、端口、token、webhook、图片上传、文件上传、Hermes session header、tool guard 身份文件，并给出可复制修复建议。",
            "QLOS Trace Viewer：把一次 QQ 消息完整链路串起来，显示 OneBot 入站、QLOS 决策、Hermes 请求、流式片段、审批、回包、文件发送、失败原因。",
            "附件流水线升级：图片、语音、文件进入统一 media index；每个附件有来源、hash、大小、MIME、OCR/ASR 状态、Hermes 是否看见、QQ 是否回传成功。",
            "QQ 口语回复实验室：记录 Hermes 原文、QLOS 清洗后文本、分隔符拆分、实际发送间隔；方便调口吻、分段和格式清洗。",
            "角色/额度表：owner、朋友、群友、临时访客独立权限和每日额度；图片、长文本、文件分析、联网、工具调用分开计量。",
        ],
    )

    add_heading(doc, "Hermi 核心功能", 1)
    add_bullets(
        doc,
        [
            "会话列表：QQ 私聊、QQ群、Hermi 新会话、工具任务、审批任务都作为可切换会话。",
            "多端同步：桌面端、手机端、PWA 共享 session_id 和 session_key；断线后可恢复正在生成的回复。",
            "Owner 控制台：服务状态、近期 trace、待审批、主动消息草稿、预算、错误日志、NapCat/QLOS/Hermes 开关。",
            "审批 inbox：把 Hermes 工具审批、UAC/管理员提示、文件写入、主动发消息等请求变成卡片，支持同意一次、拒绝、仅沙箱、永久规则。",
            "记忆审计 inbox：Hermes 想写入的记忆先进入审计箱，显示主语、来源、置信度、作用域、是否 owner 偏好。",
            "朋友模式：给朋友独立入口和额度，能聊天、总结、看文件，但不能碰 owner 文件、shell、人格和长期策略。",
            "主动行为：Hermes 可以创建“待发消息/提醒/建议”草稿，由 Hermi 或 QQ owner 审批后发出。",
            "语音和图片：手机端像对讲机一样发语音，Hermi 自动 ASR；图片进入同一附件流水线，结果可回 QQ。",
        ],
    )

    add_heading(doc, "好玩但有用", 1)
    add_bullets(
        doc,
        [
            "关系温度计：不是强数值恋爱游戏，而是轻量的互动状态，决定主动问候频率、回复热情、是否建议休息。",
            "心情/在线状态：忙碌、陪聊、项目模式、夜间低打扰；影响 Hermes 主动性和回复长度。",
            "小剧场 trace：把一次工具任务变成可读时间线，比如“收到图片 -> OCR -> 查资料 -> 生成总结 -> 发回 QQ”。",
            "回忆卡片：从日志和记忆中生成“这周我们聊了什么”，owner 可收藏、删除或转成长期记忆。",
            "好友 guest room：给朋友临时房间，自动限额、过期、只保存该朋友自己的上下文。",
            "消息风格调音台：口语/冷静/可爱/专业/短句/多句拆分，实际落地为 prompt skill + QLOS 清洗策略。",
        ],
    )

    add_heading(doc, "架构建议", 1)
    add_numbered(
        doc,
        [
            "Hermes 保持主脑；Hermi 直连 Hermes API；QLOS-Lite 只作为 QQ 接入网关。",
            "所有入口统一写 session_id、session_key、actor_id、role、permission、source，这样插件和记忆都能知道谁在说话。",
            "审批、主动消息、记忆写入都走事件流或 inbox，不要散落在 QQ 特殊逻辑里。",
            "QLOS-Lite 不变胖：新增能力优先做成 trace、doctor、queue、attachment pipeline、role table，而不是做新 Agent。",
            "未来 Hermi Gateway 可统一承接 QQ、桌面、手机、PWA、朋友访问，再转发 Hermes。",
        ],
    )

    add_heading(doc, "资料来源", 1)
    for source in SOURCES:
        doc.add_paragraph(f"{source['name']}: {source['url']}")

    doc.save(path)


def load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def ask_hermes() -> str:
    load_env_file(ROOT / "secrets.local.env")
    config = load_qlos_lite_config()
    api_key = config.hermes_api_key
    if not api_key:
        return "Hermes API key 未配置，无法调用 Hermes 生成思索。"

    source_text = "\n".join(f"- {s['name']}: {s['takeaway']} {s['url']}" for s in SOURCES)
    prompt = f"""
你是 Hermes。请基于资料和项目背景，思考 QLOS-Lite 与 Hermi 还能增加哪些好玩且实用的功能。

要求：
1. 用中文。
2. 分近期、中期、长期。
3. 偏个人 owner + QQ + 自建桌面端/手机端/PWA 场景。
4. 具体，不要空泛营销词。
5. 注意 QLOS-Lite 仍然保持薄网关，主脑仍然是 Hermes。

项目背景：
{QLOS_HERMI_CONTEXT}

调研资料：
{source_text}
"""
    payload = {
        "model": "hermes-agent",
        "messages": [
            {"role": "system", "content": "你是 owner 的长期 AI 主脑，请给出产品和工程建议。"},
            {"role": "user", "content": prompt},
        ],
        "stream": False,
    }
    request = urllib.request.Request(
        config.hermes_api_url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "X-Hermes-Session-Id": "codex-hermi-thoughts",
            "X-Hermes-Session-Key": f"qlos:qq:dm:{config.owner_qq_id or 'owner'}",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            data = json.loads(response.read().decode("utf-8", errors="replace"))
        return data["choices"][0]["message"]["content"]
    except Exception as exc:  # pragma: no cover - operational fallback
        return f"Hermes 调用失败：{exc}"


def build_hermes_doc(path: Path, hermes_text: str) -> None:
    doc = Document()
    doc.add_heading("hermes思索：QLOS-Lite 与 Hermi 可玩功能", 0)
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    add_heading(doc, "提供给 Hermes 的项目背景", 1)
    doc.add_paragraph(QLOS_HERMI_CONTEXT.strip())
    add_heading(doc, "提供给 Hermes 的资料摘要", 1)
    add_bullets(doc, [f"{s['name']}：{s['takeaway']}（{s['url']}）" for s in SOURCES])
    add_heading(doc, "Hermes 输出", 1)
    for block in hermes_text.splitlines():
        text = block.strip()
        if text:
            doc.add_paragraph(text)
    doc.save(path)


def send_docs(paths: list[Path]) -> dict[str, Any]:
    config = load_qlos_lite_config()
    client = OneBotClient(config.onebot_base_url, config.onebot_access_token, timeout=60)
    user_id = config.owner_qq_id or "1000000001"
    result: dict[str, Any] = {"sent": [], "failed": []}
    for path in paths:
        try:
            client.upload_private_file(user_id, str(path), path.name)
            result["sent"].append(str(path))
        except Exception as exc:  # pragma: no cover - depends on NapCat
            result["failed"].append({"path": str(path), "error": str(exc)})
            try:
                client.send_private_msg(user_id, f"{path.name} 已生成，但 OneBot 发文件失败：{exc}\n本地路径：{path}")
            except Exception:
                pass
    return result


def main() -> None:
    out_dir = ROOT / "docs"
    out_dir.mkdir(exist_ok=True)
    codex_path = out_dir / "codex思索.docx"
    hermes_path = out_dir / "hermes思索.docx"
    build_codex_doc(codex_path)
    hermes_text = ask_hermes()
    build_hermes_doc(hermes_path, hermes_text)
    send_result = send_docs([codex_path, hermes_path])
    print(json.dumps({"docs": [str(codex_path), str(hermes_path)], "send": send_result}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

